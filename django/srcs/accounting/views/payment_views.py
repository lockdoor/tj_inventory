import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
import datetime
from decimal import Decimal
from django.views.generic import ListView, DetailView, CreateView, UpdateView, View, TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.shortcuts import get_object_or_404, redirect
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db.models import Sum, Q
from django.http import JsonResponse
from django.utils import timezone
from accounting.models import PettyCashPayment, PettyCashAccount, PettyCashPaymentItem, PettyCashCategory
from accounting.forms.payment_form import PettyCashPaymentForm, PettyCashPaymentItemFormSet
from accounting.services.payment_service import PettyCashPaymentService


class PettyCashPaymentListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    model = PettyCashPayment
    template_name = 'accounting/payment_list.html'
    context_object_name = 'payments'
    permission_required = 'accounting.view_pettycashpayment'
    paginate_by = 20

    def get_account(self):
        return get_object_or_404(PettyCashAccount, code=self.kwargs['account_code'], is_deleted=False)

    def get_queryset(self):
        account = self.get_account()
        qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
        
        sf_list = self.request.GET.getlist('sf')
        sv_list = self.request.GET.getlist('sv')
        
        # Also fall back to general q search if present
        q = self.request.GET.get('q', '').strip()
        if q:
            qs = qs.filter(
                Q(payment_no__icontains=q) |
                Q(payee_name__icontains=q) |
                Q(note__icontains=q)
            )

        for sf, sv in zip(sf_list, sv_list):
            sv = sv.strip()
            if not sv:
                continue
            if sf == 'voucher_no':
                qs = qs.filter(payment_no__icontains=sv)
            elif sf == 'payee':
                qs = qs.filter(payee_name__icontains=sv)
            elif sf == 'gl_code':
                qs = qs.filter(items__category__code__icontains=sv)
            elif sf == 'external_pv':
                qs = qs.filter(items__external_pv_no__icontains=sv)
            elif sf == 'description':
                qs = qs.filter(Q(items__description__icontains=sv) | Q(note__icontains=sv))
            else: # 'all' or fallback
                qs = qs.filter(
                    Q(payment_no__icontains=sv) |
                    Q(payee_name__icontains=sv) |
                    Q(note__icontains=sv) |
                    Q(items__category__code__icontains=sv) |
                    Q(items__external_pv_no__icontains=sv) |
                    Q(items__description__icontains=sv)
                )
                
        return qs.distinct().select_related('account', 'created_by').prefetch_related('items')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['account'] = self.get_account()
        context['q'] = self.request.GET.get('q', '')
        
        # Build search lines list for template
        sf_list = self.request.GET.getlist('sf')
        sv_list = self.request.GET.getlist('sv')
        search_lines = []
        for sf, sv in zip(sf_list, sv_list):
            if sv.strip():
                search_lines.append({'field': sf, 'value': sv.strip()})
        
        # Always guarantee at least one search line if none exist
        if not search_lines:
            search_lines.append({'field': 'all', 'value': ''})
            
        context['search_lines'] = search_lines
        return context


class PettyCashPaymentDetailView(LoginRequiredMixin, PermissionRequiredMixin, DetailView):
    model = PettyCashPayment
    template_name = 'accounting/payment_detail.html'
    context_object_name = 'payment'
    permission_required = 'accounting.view_pettycashpayment'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['items'] = self.object.items.all().select_related('category')
        return context


class PettyCashPaymentCreateView(LoginRequiredMixin, PermissionRequiredMixin, CreateView):
    model = PettyCashPayment
    form_class = PettyCashPaymentForm
    template_name = 'accounting/payment_form.html'
    permission_required = 'accounting.add_pettycashpayment'

    def get_account(self):
        return get_object_or_404(PettyCashAccount, code=self.kwargs['account_code'], is_deleted=False)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        account = self.get_account()
        context['account'] = account
        if self.request.POST:
            context['formset'] = PettyCashPaymentItemFormSet(self.request.POST, company=account.company)
        else:
            context['formset'] = PettyCashPaymentItemFormSet(company=account.company)
        return context

    def form_valid(self, form):
        account = self.get_account()
        context = self.get_context_data()
        formset = context['formset']
        
        if formset.is_valid():
            items_data = []
            for item_form in formset:
                if item_form.cleaned_data.get('DELETE'):
                    continue
                if not item_form.cleaned_data.get('amount'):
                    continue
                items_data.append({
                    'category': item_form.cleaned_data.get('category'),
                    'description': item_form.cleaned_data.get('description', ''),
                    'amount': item_form.cleaned_data['amount'],
                    'tax': item_form.cleaned_data.get('tax'),
                    'note': item_form.cleaned_data.get('note', ''),
                    'external_pv_no': item_form.cleaned_data.get('external_pv_no', ''),
                    'rounding_adjustment': item_form.cleaned_data.get('rounding_adjustment')
                })

            try:
                PettyCashPaymentService.create_payment(
                    account=account,
                    payment_type=form.cleaned_data['payment_type'],
                    items_data=items_data,
                    payee=form.cleaned_data.get('payee'),
                    payee_name=form.cleaned_data.get('payee_name', ''),
                    payment_date=form.cleaned_data.get('payment_date'),
                    created_by=self.request.user,
                    note=form.cleaned_data.get('note', '')
                )
                messages.success(self.request, "Voucher created and balance updated successfully.")
                return redirect('accounting:payment-list', account_code=account.code)
            except ValidationError as e:
                form.add_error(None, e)
                return self.form_invalid(form)
        else:
            return self.form_invalid(form)


class PettyCashPaymentUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = PettyCashPayment
    form_class = PettyCashPaymentForm
    template_name = 'accounting/payment_form.html'
    permission_required = 'accounting.change_pettycashpayment'

    def get_queryset(self):
        return PettyCashPayment.objects.filter(is_deleted=False)

    def dispatch(self, request, *args, **kwargs):
        payment = self.get_object()
        if payment.is_posted:
            messages.error(request, "This payment is posted to Express and cannot be updated.")
            return redirect('accounting:payment-detail', pk=payment.pk)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        account = self.object.account
        context['account'] = account
        context['next'] = self.request.GET.get('next') or self.request.POST.get('next') or ''
        if self.request.POST:
            context['formset'] = PettyCashPaymentItemFormSet(self.request.POST, instance=self.object, company=account.company)
        else:
            context['formset'] = PettyCashPaymentItemFormSet(instance=self.object, company=account.company)
        return context

    def form_valid(self, form):
        account = self.object.account
        context = self.get_context_data()
        formset = context['formset']
        
        if formset.is_valid():
            items_data = []
            for item_form in formset:
                if item_form.cleaned_data.get('DELETE'):
                    continue
                if not item_form.cleaned_data.get('amount'):
                    continue
                items_data.append({
                    'category': item_form.cleaned_data.get('category') or (item_form.instance.category if item_form.instance and item_form.instance.pk else None),
                    'description': item_form.cleaned_data.get('description', ''),
                    'amount': item_form.cleaned_data['amount'],
                    'tax': item_form.cleaned_data.get('tax'),
                    'note': item_form.cleaned_data.get('note', ''),
                    'external_pv_no': item_form.cleaned_data.get('external_pv_no', ''),
                    'rounding_adjustment': item_form.cleaned_data.get('rounding_adjustment')
                })

            try:
                PettyCashPaymentService.update_payment(
                    self.object,
                    updated_by=self.request.user,
                    items_data=items_data,
                    payee=form.cleaned_data.get('payee'),
                    payee_name=form.cleaned_data.get('payee_name', ''),
                    payment_date=form.cleaned_data.get('payment_date'),
                    note=form.cleaned_data.get('note', '')
                )
                messages.success(self.request, "Voucher updated successfully.")
                next_url = self.request.GET.get('next') or self.request.POST.get('next')
                if next_url:
                    return redirect(next_url)
                return redirect('accounting:payment-list', account_code=account.code)
            except ValidationError as e:
                form.add_error(None, e)
                return self.form_invalid(form)
        else:
            return self.form_invalid(form)


class PettyCashPaymentCancelView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'accounting.delete_pettycashpayment'

    def post(self, request, pk, *args, **kwargs):
        payment = get_object_or_404(PettyCashPayment, pk=pk, is_deleted=False)
        account_code = payment.account.code
        if payment.is_posted:
            messages.error(request, "This payment is posted to Express and cannot be cancelled.")
            return redirect('accounting:payment-detail', pk=payment.pk)
        try:
            PettyCashPaymentService.cancel_payment(payment, user=request.user)
            messages.success(request, "Voucher cancelled and balance reversed successfully.")
        except ValidationError as e:
            messages.error(request, e.message)
        return redirect('accounting:payment-list', account_code=account_code)


class PettyCashPaymentTrashListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    model = PettyCashPayment
    template_name = 'accounting/payment_trash_list.html'
    context_object_name = 'payments'
    permission_required = 'accounting.delete_pettycashpayment'

    def get_account(self):
        return get_object_or_404(PettyCashAccount, code=self.kwargs['account_code'], is_deleted=False)

    def get_queryset(self):
        account = self.get_account()
        return PettyCashPayment.objects.filter(account=account, is_deleted=True).select_related('created_by')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['account'] = self.get_account()
        return context


class PettyCashPaymentSummaryView(LoginRequiredMixin, PermissionRequiredMixin, TemplateView):
    template_name = 'accounting/payment_summary.html'
    permission_required = 'accounting.change_pettycashpayment'

    def get_account(self):
        return get_object_or_404(PettyCashAccount, code=self.kwargs['account_code'], is_deleted=False)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        account = self.get_account()
        context['account'] = account

        # Fetch replenishments to build rounds
        replenishments = PettyCashPayment.objects.filter(
            account=account,
            payment_type='replenishment',
            is_deleted=False
        ).order_by('-payment_date', '-id')

        # Check Active (Unreplenished) Round
        latest_rep = replenishments.first()
        active_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
        if latest_rep:
            active_qs = active_qs.filter(id__gt=latest_rep.id)

        # Build dropdown options
        rounds = []
        if active_qs.exists():
            rounds.append({
                'id': 'active',
                'name': 'Active (Unreplenished) Round'
            })
        for rep in replenishments:
            formatted_date = rep.payment_date.strftime('%Y-%m-%d') if rep.payment_date else ''
            rounds.append({
                'id': str(rep.id),
                'name': f"Replenishment {rep.payment_no} ({formatted_date})"
            })
        context['rounds'] = rounds

        # Determine selected round
        round_id = self.request.GET.get('round_id')
        if not round_id and rounds:
            round_id = rounds[0]['id']
        context['selected_round_id'] = round_id

        # Calculate prev/next round navigation
        prev_round = None
        next_round = None
        current_idx = None
        for i, rnd in enumerate(rounds):
            if rnd['id'] == round_id:
                current_idx = i
                break
        
        if current_idx is not None:
            if current_idx < len(rounds) - 1:
                prev_round = rounds[current_idx + 1]  # Older
            if current_idx > 0:
                next_round = rounds[current_idx - 1]  # Newer
        
        context['prev_round'] = prev_round
        context['next_round'] = next_round

        selected_rep = None
        is_active_round = False

        if round_id == 'active':
            is_active_round = True
            payments_qs = active_qs
        elif round_id:
            try:
                selected_rep = replenishments.get(pk=int(round_id))
                prev_rep = replenishments.filter(id__lt=selected_rep.id).order_by('-id').first()
                payments_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
                if prev_rep:
                    payments_qs = payments_qs.filter(id__gt=prev_rep.id, id__lte=selected_rep.id)
                else:
                    payments_qs = payments_qs.filter(id__lte=selected_rep.id)
            except (ValueError, PettyCashPayment.DoesNotExist):
                payments_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
        # Apply advanced multi-condition search filtering on payments_qs
        sf_list = self.request.GET.getlist('sf')
        sv_list = self.request.GET.getlist('sv')
        
        q = self.request.GET.get('q', '').strip()
        if q:
            payments_qs = payments_qs.filter(
                Q(payment_no__icontains=q) |
                Q(payee_name__icontains=q) |
                Q(note__icontains=q)
            )

        for sf, sv in zip(sf_list, sv_list):
            sv = sv.strip()
            if not sv:
                continue
            if sf == 'voucher_no':
                payments_qs = payments_qs.filter(payment_no__icontains=sv)
            elif sf == 'payee':
                payments_qs = payments_qs.filter(payee_name__icontains=sv)
            elif sf == 'gl_code':
                payments_qs = payments_qs.filter(items__category__code__icontains=sv)
            elif sf == 'external_pv':
                payments_qs = payments_qs.filter(items__external_pv_no__icontains=sv)
            elif sf == 'description':
                payments_qs = payments_qs.filter(Q(items__description__icontains=sv) | Q(note__icontains=sv))
            else: # 'all' or fallback
                payments_qs = payments_qs.filter(
                    Q(payment_no__icontains=sv) |
                    Q(payee_name__icontains=sv) |
                    Q(note__icontains=sv) |
                    Q(items__category__code__icontains=sv) |
                    Q(items__external_pv_no__icontains=sv) |
                    Q(items__description__icontains=sv)
                )

        payments_qs = payments_qs.distinct()

        # Build search lines list for template
        search_lines = []
        for sf, sv in zip(sf_list, sv_list):
            if sv.strip():
                search_lines.append({'field': sf, 'value': sv.strip()})
        
        if not search_lines:
            search_lines.append({'field': 'all', 'value': ''})
            
        context['search_lines'] = search_lines
        context['q'] = q

        items = PettyCashPaymentItem.objects.filter(payment__in=payments_qs)
        
        round_items = items.exclude(payment__payment_type='replenishment')

        # Distinguish actual PV items and normal items
        actual_pv_items = round_items.exclude(external_pv_no='').exclude(external_pv_no__isnull=True).order_by('external_pv_no', 'id')
        normal_items = round_items.filter(Q(external_pv_no='') | Q(external_pv_no__isnull=True))

        # In-memory aggregation of normal items with VAT extraction and rounding adjustment deduction
        category_sums_dict = {}
        total_vat = Decimal('0.00')
        total_rounding = Decimal('0.00')
        unallocated_sum = Decimal('0.00')

        # Prefetch payment to avoid N+1 queries on item.payment
        normal_items = normal_items.select_related('payment')

        vat_items_list = []
        rounding_items_list = []
        for item in normal_items:
            tax_amount = item.tax or Decimal('0.00')
            item_rounding = item.rounding_adjustment or Decimal('0.00')

            net_amount = item.amount - tax_amount - item_rounding

            if tax_amount > Decimal('0.00'):
                vat_code = account.vat_category_code or '1155-00'
                vat_cat = PettyCashCategory.objects.filter(code=vat_code, company=account.company, is_deleted=False).first()
                vat_name = vat_cat.name if vat_cat else "ภาษีซื้อ-ยังไม่ถึงกำหนด"
                payee = item.payment.payee_name or (item.payment.created_by.get_full_name() if item.payment.created_by else '') or str(item.payment.created_by or '')
                desc_str = f"VAT: {payee} - {item.description}" if payee and item.description else (payee or item.description or 'Input VAT')
                vat_items_list.append({
                    'category__code': vat_code,
                    'category__name': f"{vat_name} ({desc_str})",
                    'total': tax_amount
                })

            if item_rounding != Decimal('0.00'):
                rounding_code = account.rounding_category_code or '4200-07'
                rounding_cat = PettyCashCategory.objects.filter(code=rounding_code, company=account.company, is_deleted=False).first()
                rounding_name = rounding_cat.name if rounding_cat else "รายได้-อื่นๆ"
                payee = item.payment.payee_name or (item.payment.created_by.get_full_name() if item.payment.created_by else '') or str(item.payment.created_by or '')
                desc_str = f"Rounding: {payee} - {item.description}" if payee and item.description else (payee or item.description or 'Rounding Adjustment')
                rounding_items_list.append({
                    'category__code': rounding_code,
                    'category__name': f"{rounding_name} ({desc_str})",
                    'total': item_rounding
                })

            if item.category:
                code = item.category.code
                name = item.category.name
                if code not in category_sums_dict:
                    category_sums_dict[code] = {
                        'category__code': code,
                        'category__name': name,
                        'total': Decimal('0.00')
                    }
                category_sums_dict[code]['total'] += net_amount
            else:
                unallocated_sum += net_amount

        # Convert to list and sort normal categories by code
        category_sums_list = sorted(category_sums_dict.values(), key=lambda x: x['category__code'])

        # Append individual VAT records
        category_sums_list.extend(vat_items_list)

        # Append individual rounding records
        category_sums_list.extend(rounding_items_list)

        # Append unallocated row if it exists
        if unallocated_sum > Decimal('0.00'):
            category_sums_list.append({
                'category__code': None,
                'category__name': "Pending category allocation",
                'total': unallocated_sum
            })

        # Append individual actual PV records
        for item in actual_pv_items:
            payment = item.payment
            payee = payment.payee_name or (payment.payee.full_name if payment.payee else '') or str(payment.payee or '')
            desc_str = f"{payee} - {item.description}" if payee and item.description else (payee or item.description or 'External PV')
            category_sums_list.append({
                'category__code': f"PV: {item.external_pv_no}",
                'category__name': desc_str,
                'total': item.amount
            })

        # Compute unallocated count, excluding items belonging to replenishment or actual PV items
        unallocated_count = items.filter(category__isnull=True).exclude(payment__payment_type='replenishment').exclude(
            ~Q(external_pv_no='') & Q(external_pv_no__isnull=False)
        ).count()

        context['payments'] = payments_qs
        context['category_sums'] = category_sums_list
        context['total_spent'] = sum(row['total'] for row in category_sums_list)
        context['unallocated_count'] = unallocated_count
        context['unposted_payments'] = payments_qs.filter(is_posted=False)
        context['posted_payments'] = payments_qs.filter(is_posted=True)
        context['is_active_round'] = is_active_round
        context['selected_rep'] = selected_rep

        return context

    def post(self, request, *args, **kwargs):
        account = self.get_account()
        round_id = request.POST.get('round_id')

        if round_id == 'active':
            messages.error(request, "You cannot lock the active round until a replenishment record is created.")
            return redirect(f"{request.path}?round_id=active")

        selected_rep = get_object_or_404(PettyCashPayment, pk=int(round_id), account=account)
        replenishments = PettyCashPayment.objects.filter(
            account=account,
            payment_type='replenishment',
            is_deleted=False
        ).order_by('-id')
        prev_rep = replenishments.filter(id__lt=selected_rep.id).first()

        payments_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
        if prev_rep:
            payments_qs = payments_qs.filter(id__gt=prev_rep.id, id__lte=selected_rep.id)
        else:
            payments_qs = payments_qs.filter(id__lte=selected_rep.id)

        unposted_payments = payments_qs.filter(is_posted=False)

        if not unposted_payments.exists():
            messages.warning(request, "No unposted vouchers found for the selected round.")
            return redirect(f"{request.path}?round_id={round_id}")

        try:
            PettyCashPaymentService.mark_payments_as_posted(unposted_payments, user=request.user)
            messages.success(request, f"Successfully marked {len(unposted_payments)} vouchers as posted to Express.")
        except ValidationError as e:
            messages.error(request, e.message)

        return redirect(f"{request.path}?round_id={round_id}")


class PettyCashCategorySearchAPIView(LoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        q = request.GET.get('q', '').strip()
        company_id = request.GET.get('company_id')
        payment_id = request.GET.get('payment_id')
        
        qs = PettyCashCategory.objects.filter(is_deleted=False)
        if payment_id:
            payment = get_object_or_404(PettyCashPayment, pk=payment_id)
            qs = qs.filter(company=payment.account.company)
        elif company_id:
            qs = qs.filter(company_id=company_id)
            
        if q:
            qs = qs.filter(Q(code__icontains=q) | Q(name__icontains=q))
            
        # Limit to top 15 results for performance
        results = [
            {'id': cat.id, 'code': cat.code, 'name': cat.name}
            for cat in qs[:15]
        ]
        return JsonResponse({'results': results})


class PettyCashPaymentAllocateAPIView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'accounting.change_pettycashpayment'
    
    def post(self, request, pk, *args, **kwargs):
        item = get_object_or_404(PettyCashPaymentItem, pk=pk)
        if item.payment.is_posted:
            return JsonResponse({'error': 'Cannot modify posted vouchers.'}, status=400)
            
        import json
        try:
            data = json.loads(request.body)
        except json.JSONDecodeError:
            return JsonResponse({'error': 'Invalid JSON.'}, status=400)
            
        category_id = data.get('category_id')
        external_pv_no = data.get('external_pv_no')
        
        if not category_id and not external_pv_no:
            return JsonResponse({'error': 'Either Category or External PV number is required.'}, status=400)
            
        if category_id:
            category = get_object_or_404(PettyCashCategory, pk=category_id, is_deleted=False)
            if category.company != item.payment.account.company:
                return JsonResponse({'error': 'Category company mismatch.'}, status=400)
                
            item.category = category
            item.external_pv_no = ''
            item.save()
        else:
            item.category = None
            item.external_pv_no = external_pv_no.strip()
            item.save()
            
        return JsonResponse({'success': True})


class PettyCashPaymentSummaryExportView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'accounting.view_pettycashpayment'

    def get_account(self, account_code):
        return get_object_or_404(PettyCashAccount, code=account_code, is_deleted=False)

    def get(self, request, account_code, *args, **kwargs):
        account = self.get_account(account_code)

        # Fetch replenishments to resolve rounds
        replenishments = PettyCashPayment.objects.filter(
            account=account,
            payment_type='replenishment',
            is_deleted=False
        ).order_by('-payment_date', '-id')

        latest_rep = replenishments.first()
        active_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
        if latest_rep:
            active_qs = active_qs.filter(id__gt=latest_rep.id)

        rounds = []
        if active_qs.exists():
            rounds.append({'id': 'active', 'name': 'Active (Unreplenished) Round'})
        for rep in replenishments:
            formatted_date = rep.payment_date.strftime('%Y-%m-%d') if rep.payment_date else ''
            rounds.append({
                'id': str(rep.id),
                'name': f"Replenishment {rep.payment_no} ({formatted_date})"
            })

        round_id = request.GET.get('round_id')
        if not round_id and rounds:
            round_id = rounds[0]['id']

        selected_rep = None
        if round_id == 'active':
            payments_qs = active_qs
        elif round_id:
            try:
                selected_rep = replenishments.get(pk=int(round_id))
                prev_rep = replenishments.filter(id__lt=selected_rep.id).order_by('-id').first()
                payments_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
                if prev_rep:
                    payments_qs = payments_qs.filter(id__gt=prev_rep.id, id__lte=selected_rep.id)
                else:
                    payments_qs = payments_qs.filter(id__lte=selected_rep.id)
            except (ValueError, PettyCashPayment.DoesNotExist):
                payments_qs = PettyCashPayment.objects.filter(account=account, is_deleted=False)
        else:
            payments_qs = PettyCashPayment.objects.none()

        # Apply multi-condition search filtering if any
        sf_list = request.GET.getlist('sf')
        sv_list = request.GET.getlist('sv')
        q = request.GET.get('q', '').strip()
        if q:
            payments_qs = payments_qs.filter(
                Q(payment_no__icontains=q) |
                Q(payee_name__icontains=q) |
                Q(note__icontains=q)
            )

        for sf, sv in zip(sf_list, sv_list):
            sv = sv.strip()
            if not sv:
                continue
            if sf == 'voucher_no':
                payments_qs = payments_qs.filter(payment_no__icontains=sv)
            elif sf == 'payee':
                payments_qs = payments_qs.filter(payee_name__icontains=sv)
            elif sf == 'gl_code':
                payments_qs = payments_qs.filter(items__category__code__icontains=sv)
            elif sf == 'external_pv':
                payments_qs = payments_qs.filter(items__external_pv_no__icontains=sv)
            elif sf == 'description':
                payments_qs = payments_qs.filter(Q(items__description__icontains=sv) | Q(note__icontains=sv))
            else:
                payments_qs = payments_qs.filter(
                    Q(payment_no__icontains=sv) |
                    Q(payee_name__icontains=sv) |
                    Q(note__icontains=sv) |
                    Q(items__category__code__icontains=sv) |
                    Q(items__external_pv_no__icontains=sv) |
                    Q(items__description__icontains=sv)
                )

        payments_qs = payments_qs.distinct()

        items = PettyCashPaymentItem.objects.filter(
            payment__in=payments_qs
        ).exclude(
            payment__payment_type='replenishment'
        ).select_related('payment', 'payment__payee').order_by('payment__payment_date', 'payment__id', 'id')

        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ใบเบิกเงินสดย่อย"

        # Styles
        font_company = Font(name='Arial', size=14, bold=True)
        font_title = Font(name='Arial', size=13, bold=True)
        font_date = Font(name='Arial', size=10, italic=True)
        font_header = Font(name='Arial', size=10, bold=True)
        font_data = Font(name='Arial', size=10)
        font_total = Font(name='Arial', size=10, bold=True)

        header_fill = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
        total_fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')

        thin_side = Side(style='thin', color='C0C0C0')
        double_side = Side(style='double', color='000000')

        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        total_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_side)

        # 3 Headers out of table
        company_name = account.company.name if account.company else "องค์กร"
        ws['A1'] = company_name
        ws['A1'].font = font_company
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A1:F1')

        ws['A2'] = "ใบเบิกเงินสดย่อย"
        ws['A2'].font = font_title
        ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A2:F2')

        if selected_rep and selected_rep.payment_date:
            rep_date_str = selected_rep.payment_date.strftime('%d/%m/%Y')
            ws['A3'] = f"วันที่เบิกชดเชย: {rep_date_str}"
        elif round_id == 'active':
            ws['A3'] = "วันที่เบิกชดเชย: รอบปัจจุบัน (ยังไม่ได้เบิกชดเชย)"
        else:
            ws['A3'] = "วันที่เบิกชดเชย: -"
        ws['A3'].font = font_date
        ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A3:F3')

        # Table Headers at Row 5
        headers = ["ลำดับ", "วันที่", "จ่ายให้", "รายการ", "ภาษี", "ยอดเงิน"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.border = cell_border
            if col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in [5, 6]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        current_row = 6
        num_items = items.count()

        for idx, item in enumerate(items, start=1):
            date_str = item.payment.payment_date.strftime('%d/%m/%Y') if item.payment.payment_date else ''
            payee = item.payment.payee_name or (item.payment.payee.full_name if item.payment.payee else '') or str(item.payment.payee or '')
            desc = item.description or ''
            tax_val = float(item.tax) if item.tax is not None else 0.0
            amount_val = float(item.amount) if item.amount is not None else 0.0

            cA = ws.cell(row=current_row, column=1, value=idx)
            cA.alignment = Alignment(horizontal='center', vertical='center')

            cB = ws.cell(row=current_row, column=2, value=date_str)
            cB.alignment = Alignment(horizontal='center', vertical='center')

            cC = ws.cell(row=current_row, column=3, value=payee)
            cC.alignment = Alignment(horizontal='left', vertical='center')

            cD = ws.cell(row=current_row, column=4, value=desc)
            cD.alignment = Alignment(horizontal='left', vertical='center')

            cE = ws.cell(row=current_row, column=5, value=tax_val)
            cE.alignment = Alignment(horizontal='right', vertical='center')
            cE.number_format = '#,##0.00'

            cF = ws.cell(row=current_row, column=6, value=amount_val)
            cF.alignment = Alignment(horizontal='right', vertical='center')
            cF.number_format = '#,##0.00'

            for c in [cA, cB, cC, cD, cE, cF]:
                c.font = font_data
                c.border = cell_border

            current_row += 1

        # Summary Row
        summary_row = current_row
        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
        c_label = ws.cell(row=summary_row, column=1, value="รวม")
        c_label.font = font_total
        c_label.alignment = Alignment(horizontal='right', vertical='center')

        for col_idx in range(1, 5):
            cell = ws.cell(row=summary_row, column=col_idx)
            cell.border = total_border
            cell.fill = total_fill

        c_sum_tax = ws.cell(row=summary_row, column=5)
        c_sum_amount = ws.cell(row=summary_row, column=6)

        if num_items > 0:
            c_sum_tax.value = f"=SUM(E6:E{summary_row - 1})"
            c_sum_amount.value = f"=SUM(F6:F{summary_row - 1})"
        else:
            c_sum_tax.value = 0.00
            c_sum_amount.value = 0.00

        for cell in [c_sum_tax, c_sum_amount]:
            cell.font = font_total
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0.00'
            cell.border = total_border
            cell.fill = total_fill

        # Column widths
        column_widths = {
            'A': 8,
            'B': 13,
            'C': 24,
            'D': 36,
            'E': 14,
            'F': 16,
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        # Row heights
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[5].height = 22
        for r in range(6, current_row + 1):
            ws.row_dimensions[r].height = 20

        # Print / Page Setup for A4
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        clean_round = "".join(c for c in (round_id or 'all') if c.isalnum() or c in ('-', '_'))
        filename = f"petty_cash_{account.code}_{clean_round}.xlsx"

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
