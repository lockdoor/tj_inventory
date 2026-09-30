import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from decimal import Decimal


class PettyCashExcelExportService:
    """
    Service responsible for generating print-ready Excel spreadsheets
    for petty cash replenishment documents and category summaries.
    """

    # Shared Styles
    FONT_COMPANY = Font(name='Arial', size=14, bold=True)
    FONT_TITLE = Font(name='Arial', size=13, bold=True)
    FONT_DATE = Font(name='Arial', size=10, italic=True)
    FONT_HEADER = Font(name='Arial', size=10, bold=True)
    FONT_DATA = Font(name='Arial', size=10)
    FONT_TOTAL = Font(name='Arial', size=10, bold=True)
    FONT_SIGN = Font(name='Arial', size=9)

    HEADER_FILL = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    TOTAL_FILL = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')

    THIN_SIDE = Side(style='thin', color='C0C0C0')
    DOUBLE_SIDE = Side(style='double', color='000000')

    CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
    TOTAL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=DOUBLE_SIDE)

    @classmethod
    def apply_a4_page_setup(cls, ws):
        """Configure sheet for A4 portrait printing without horizontal clipping."""
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    @classmethod
    def generate_replenishment_vouchers_excel(cls, account, round_id, selected_rep, items):
        """
        Generates 'ใบเบิกเงินสดย่อย' itemized voucher requisition spreadsheet.
        Returns: (BytesIO buffer, filename)
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ใบเบิกเงินสดย่อย"

        # Headers above table
        company_name = account.company.name if account.company else "องค์กร"
        ws['A1'] = company_name
        ws['A1'].font = cls.FONT_COMPANY
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A1:F1')

        ws['A2'] = "ใบเบิกเงินสดย่อย"
        ws['A2'].font = cls.FONT_TITLE
        ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A2:F2')

        if selected_rep and selected_rep.payment_date:
            rep_date_str = selected_rep.payment_date.strftime('%d/%m/%Y')
            ws['A3'] = f"วันที่เบิกชดเชย: {rep_date_str}"
        elif round_id == 'active':
            ws['A3'] = "วันที่เบิกชดเชย: รอบปัจจุบัน (ยังไม่ได้เบิกชดเชย)"
        else:
            ws['A3'] = "วันที่เบิกชดเชย: -"
        ws['A3'].font = cls.FONT_DATE
        ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A3:F3')

        # Table Headers at Row 5
        headers = ["ลำดับ", "วันที่", "จ่ายให้", "รายการ", "ภาษี", "ยอดเงิน"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.HEADER_FILL
            cell.border = cls.CELL_BORDER
            if col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in [5, 6]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        current_row = 6
        num_items = items.count() if hasattr(items, 'count') else len(items)

        for idx, item in enumerate(items, start=1):
            date_str = item.payment.payment_date.strftime('%d/%m/%Y') if item.payment.payment_date else ''
            payee = item.payment.payee_name or (item.payment.payee.full_name if item.payment.payee else '') or str(item.payment.payee or '')
            desc = item.description or ''
            tax_val = float(item.tax) if item.tax is not None else 0.0
            amount_val = float(item.amount) if item.amount is not None else 0.0

            cA = ws.cell(row=current_row, column=1, value=idx)
            cA.alignment = Alignment(horizontal='center', vertical='center')

            cB = ws.cell(row=current_row, column=2, value=date_str)
            cB.alignment = Alignment(horizontal='center', vertical='center')

            cC = ws.cell(row=current_row, column=3, value=payee)
            cC.alignment = Alignment(horizontal='left', vertical='center')

            cD = ws.cell(row=current_row, column=4, value=desc)
            cD.alignment = Alignment(horizontal='left', vertical='center')

            cE = ws.cell(row=current_row, column=5, value=tax_val)
            cE.alignment = Alignment(horizontal='right', vertical='center')
            cE.number_format = '#,##0.00'

            cF = ws.cell(row=current_row, column=6, value=amount_val)
            cF.alignment = Alignment(horizontal='right', vertical='center')
            cF.number_format = '#,##0.00'

            for c in [cA, cB, cC, cD, cE, cF]:
                c.font = cls.FONT_DATA
                c.border = cls.CELL_BORDER

            current_row += 1

        # Summary Row
        summary_row = current_row
        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
        c_label = ws.cell(row=summary_row, column=1, value="รวม")
        c_label.font = cls.FONT_TOTAL
        c_label.alignment = Alignment(horizontal='right', vertical='center')

        for col_idx in range(1, 5):
            cell = ws.cell(row=summary_row, column=col_idx)
            cell.border = cls.TOTAL_BORDER
            cell.fill = cls.TOTAL_FILL

        c_sum_tax = ws.cell(row=summary_row, column=5)
        c_sum_amount = ws.cell(row=summary_row, column=6)

        if num_items > 0:
            c_sum_tax.value = f"=SUM(E6:E{summary_row - 1})"
            c_sum_amount.value = f"=SUM(F6:F{summary_row - 1})"
        else:
            c_sum_tax.value = 0.00
            c_sum_amount.value = 0.00

        for cell in [c_sum_tax, c_sum_amount]:
            cell.font = cls.FONT_TOTAL
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0.00'
            cell.border = cls.TOTAL_BORDER
            cell.fill = cls.TOTAL_FILL

        # Column widths
        column_widths = {
            'A': 8,
            'B': 13,
            'C': 24,
            'D': 36,
            'E': 14,
            'F': 16,
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        # Row heights
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[5].height = 22
        for r in range(6, current_row + 1):
            ws.row_dimensions[r].height = 20

        cls.apply_a4_page_setup(ws)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        clean_round = "".join(c for c in (round_id or 'all') if c.isalnum() or c in ('-', '_'))
        filename = f"petty_cash_{account.code}_{clean_round}.xlsx"

        return buffer, filename

    @classmethod
    def generate_category_summary_excel(cls, account, round_id, selected_rep, category_sums_list):
        """
        Generates 'ใบสรุปค่าใช้จ่ายเงินสดย่อยตามหมวดบัญชี' category aggregations spreadsheet.
        Returns: (BytesIO buffer, filename)
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "สรุปตามหมวดบัญชี"

        # Headers above table
        company_name = account.company.name if account.company else "องค์กร"
        ws['A1'] = company_name
        ws['A1'].font = cls.FONT_COMPANY
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A1:C1')

        ws['A2'] = "ใบสรุปค่าใช้จ่ายเงินสดย่อยตามหมวดบัญชี"
        ws['A2'].font = cls.FONT_TITLE
        ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A2:C2')

        # Resolve PV from first line of replenishment (Express accounting PV) if provided
        first_rep_item = selected_rep.items.first() if selected_rep and hasattr(selected_rep, 'items') else None
        external_pv = (first_rep_item.external_pv_no.strip() if first_rep_item and first_rep_item.external_pv_no else '') if selected_rep else ''
        pv_suffix = f" ({external_pv})" if external_pv else ""

        if selected_rep and selected_rep.payment_date:
            rep_date_str = selected_rep.payment_date.strftime('%d/%m/%Y')
            ws['A3'] = f"วงเงินสดย่อย: {account.name} ({account.code}) | วันที่เบิกชดเชย: {rep_date_str}{pv_suffix}"
        elif round_id == 'active':
            ws['A3'] = f"วงเงินสดย่อย: {account.name} ({account.code}) | วันที่เบิกชดเชย: รอบปัจจุบัน (ยังไม่ได้เบิกชดเชย)"
        else:
            ws['A3'] = f"วงเงินสดย่อย: {account.name} ({account.code}) | วันที่เบิกชดเชย: -{pv_suffix}"
        ws['A3'].font = cls.FONT_DATE
        ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A3:C3')

        # Table Headers at Row 5
        headers = ["รหัสหมวดบัญชี", "ชื่อหมวดหมู่ / รายการ", "ยอดเงิน (บาท)"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.HEADER_FILL
            cell.border = cls.CELL_BORDER
            if col_idx == 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        current_row = 6
        num_items = len(category_sums_list)

        for row in category_sums_list:
            cat_code = row.get('category__code') or 'UNALLOCATED'
            # Must show only category name, not extended describe (e.g. 1155-00 must show 'ภาษีซื้อ-ยังไม่ถึงกำหนด' only)
            cat_name = row.get('category_base_name')
            if not cat_name:
                raw_name = row.get('category__name') or 'Pending category allocation'
                if ' (' in raw_name and ('VAT:' in raw_name or 'Rounding:' in raw_name or cat_code in ('1155-00', '4200-07')):
                    cat_name = raw_name.split(' (', 1)[0]
                else:
                    cat_name = raw_name
            total_val = float(row.get('total') or 0)

            cA = ws.cell(row=current_row, column=1, value=cat_code)
            cA.alignment = Alignment(horizontal='center', vertical='center')

            cB = ws.cell(row=current_row, column=2, value=cat_name)
            cB.alignment = Alignment(horizontal='left', vertical='center')

            cC = ws.cell(row=current_row, column=3, value=total_val)
            cC.alignment = Alignment(horizontal='right', vertical='center')
            cC.number_format = '#,##0.00'

            for c in [cA, cB, cC]:
                c.font = cls.FONT_DATA
                c.border = cls.CELL_BORDER

            current_row += 1

        # Summary Row
        summary_row = current_row
        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=2)
        c_label = ws.cell(row=summary_row, column=1, value="รวมทั้งสิ้น")
        c_label.font = cls.FONT_TOTAL
        c_label.alignment = Alignment(horizontal='right', vertical='center')

        for col_idx in [1, 2]:
            cell = ws.cell(row=summary_row, column=col_idx)
            cell.border = cls.TOTAL_BORDER
            cell.fill = cls.TOTAL_FILL

        c_sum_amount = ws.cell(row=summary_row, column=3)
        if num_items > 0:
            c_sum_amount.value = f"=SUM(C6:C{summary_row - 1})"
        else:
            c_sum_amount.value = 0.00
        c_sum_amount.font = cls.FONT_TOTAL
        c_sum_amount.alignment = Alignment(horizontal='right', vertical='center')
        c_sum_amount.number_format = '#,##0.00'
        c_sum_amount.border = cls.TOTAL_BORDER
        c_sum_amount.fill = cls.TOTAL_FILL


        # Column widths
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 46
        ws.column_dimensions['C'].width = 20

        # Row heights
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[5].height = 22
        for r in range(6, current_row + 1):
            ws.row_dimensions[r].height = 20

        cls.apply_a4_page_setup(ws)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        clean_round = "".join(c for c in (round_id or 'all') if c.isalnum() or c in ('-', '_'))
        filename = f"petty_cash_categories_{account.code}_{clean_round}.xlsx"

        return buffer, filename
